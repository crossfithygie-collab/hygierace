#!/usr/bin/env python3
"""Construit race4-juges.html (page NON listée dans le menu, noindex) depuis le
planning Excel des juges. Relancer après chaque nouvelle version du fichier :

    python3 build_juges.py ~/Downloads/Hygie_Race_4_planning_juges_v4.xlsx
"""
import html
import sys
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1
           else Path.home() / "Downloads/Hygie_Race_4_planning_juges_v4.xlsx")
OUT = Path(__file__).resolve().parent / "race4-juges.html"
E = html.escape


def hhmm(v):
    s = str(v or "").strip()
    return s[:5] if len(s) >= 5 and s[2] == ":" else s


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    pl = wb["planning - volunteers"]
    lignes = [[c for c in r] for r in pl.iter_rows(values_only=True)]
    entetes = [str(c) for c in lignes[0] if c]          # Hygie Race, HEAT 1..7
    heats = entetes[1:]
    par_cle = {str(r[0]).strip(): r[1:len(entetes)] for r in lignes[1:] if r and r[0]}
    warm = [hhmm(x) for x in par_cle.get("WARM-UP", [])]
    call = [hhmm(x) for x in par_cle.get("CALL ROOM", [])]
    creneau = [str(x or "").replace(":00 ", " ").replace(":00", "") for x in par_cle.get("START-END", [])]
    divisions = [str(x or "") for x in par_cle.get("DIVISIONS", [])]
    lanes = [(k, v) for k, v in par_cle.items() if k.upper().startswith("LANE")]

    # Récap : remarques par juge
    remarques = {}
    for r in wb["Récap juges"].iter_rows(min_row=2, values_only=True):
        if not (r and r[0]):
            continue
        # « Modifié (avant : …) » est une note de travail interne : pas pour les juges.
        utiles = [m.strip() for m in str(r[3] or "").split("·")
                  if m.strip() and not m.strip().lower().startswith("modifié")]
        remarques[str(r[0]).strip()] = " · ".join(utiles)

    # Feuilles individuelles : un onglet par juge
    ignore = {"planning - volunteers", "Récap juges"} | {f"HEAT {i}" for i in range(1, 12)}
    juges = []
    for nom in wb.sheetnames:
        if nom in ignore:
            continue
        ws = wb[nom]
        rows = [[str(c or "").strip() for c in r] for r in ws.iter_rows(values_only=True)]
        # La ligne d'en-tête du fichier commence aussi par « Heat » : on l'écarte.
        creneaux = [r for r in rows if r and len(r) > 4
                    and r[0].upper().startswith("HEAT ") and r[1].lower() != "lane"]
        juges.append({"nom": nom, "creneaux": creneaux, "remarque": remarques.get(nom, "")})
    juges.sort(key=lambda j: j["nom"].lower())

    # ---- tableau général -------------------------------------------------
    th = "".join(f"<th>{E(h)}</th>" for h in heats)
    def ligne(titre, vals, cls=""):
        return (f'<tr class="{cls}"><th scope="row">{E(titre)}</th>'
                + "".join(f"<td>{E(str(v or '—'))}</td>" for v in vals[:len(heats)]) + "</tr>")
    corps = (ligne("Division", divisions, "div")
             + ligne("Warm-up", warm) + ligne("Call room", call) + ligne("Heat", creneau, "heat")
             + "".join(ligne(k.title(), [str(x or "—") for x in v], "lane") for k, v in lanes))

    # ---- fiches par juge -------------------------------------------------
    cartes = []
    for j in juges:
        li = "".join(
            f'<li><b>{E(c[0])}</b> <span class="lane">{E(c[1])}</span>'
            f'<span class="quand">{E(hhmm(c[3]))} – {E(hhmm(c[4]))}</span>'
            f'<span class="divi">{E(c[2])}</span></li>' for c in j["creneaux"])
        rq = f'<p class="rq">{E(j["remarque"])}</p>' if j["remarque"] else ""
        cartes.append(f'<article class="juge" data-nom="{E(j["nom"].lower())}">'
                      f'<h3>{E(j["nom"])}</h3><ul>{li}</ul>{rq}</article>')

    page = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Juges · Hygie Race 4</title>
<meta name="robots" content="noindex, nofollow">
<meta name="theme-color" content="#1800AD">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Archivo:ital,wght@0,400;0,600;0,700;0,800;0,900;1,800;1,900&family=Permanent+Marker&display=swap" rel="stylesheet">
<link rel="stylesheet" href="race4.css?v=4">
<style>
  .juges-wrap{{max-width:980px;margin:0 auto;padding:34px 20px 90px;position:relative;z-index:2;}}
  .juges-wrap h1{{font-size:clamp(34px,9vw,54px);text-align:center;}}
  .intro{{text-align:center;color:var(--soft);font-size:14.5px;line-height:1.6;margin:18px auto 0;max-width:620px;}}
  .intro b{{color:#fff;}}
  h2{{margin:44px 0 14px;font-weight:900;font-style:italic;font-size:15px;letter-spacing:2px;text-transform:uppercase;}}
  .scroll{{overflow-x:auto;-webkit-overflow-scrolling:touch;border:1.5px solid var(--hair);border-radius:12px;background:var(--fill);}}
  table{{border-collapse:collapse;width:100%;min-width:720px;font-size:13px;}}
  th,td{{padding:9px 11px;text-align:left;border-bottom:1px solid var(--hair);white-space:nowrap;}}
  thead th{{font-size:11px;letter-spacing:1px;text-transform:uppercase;color:var(--soft);}}
  th[scope=row]{{font-size:11px;letter-spacing:.6px;text-transform:uppercase;color:var(--soft2);font-weight:800;}}
  tr.div td{{font-weight:900;font-style:italic;}}
  tr.heat td{{font-weight:800;}}
  tr.lane td{{color:#fff;}}
  tr:last-child th,tr:last-child td{{border-bottom:0;}}
  .recherche{{display:block;width:100%;max-width:340px;margin:0 auto;padding:13px 16px;border-radius:999px;
    border:1.5px solid var(--hair);background:var(--fill);color:#fff;font:inherit;font-size:15px;text-align:center;}}
  .recherche::placeholder{{color:var(--soft2);}}
  .grille{{display:grid;gap:12px;grid-template-columns:repeat(auto-fill,minmax(270px,1fr));margin-top:16px;}}
  .juge{{border:1.5px solid var(--hair);border-radius:12px;padding:15px 17px;background:var(--fill);}}
  .juge h3{{margin:0 0 10px;font-weight:900;font-style:italic;font-size:17px;text-transform:uppercase;}}
  .juge ul{{list-style:none;margin:0;padding:0;}}
  .juge li{{display:flex;flex-wrap:wrap;gap:8px;align-items:baseline;padding:7px 0;border-top:1px solid var(--hair);font-size:13.5px;}}
  .juge li b{{font-weight:900;}}
  .lane{{font-size:11px;letter-spacing:.5px;text-transform:uppercase;border:1px solid var(--hair);border-radius:999px;padding:2px 8px;}}
  .quand{{color:#fff;font-weight:700;}}
  .divi{{width:100%;color:var(--soft2);font-size:11.5px;letter-spacing:.4px;text-transform:uppercase;}}
  .rq{{margin:10px 0 0;color:var(--soft);font-size:12.5px;line-height:1.5;}}
  .juge.cache{{display:none;}}
  .rappels{{margin-top:34px;border-left:3px solid #fff;padding:4px 0 4px 16px;color:var(--soft);font-size:14px;line-height:1.75;}}
  .rappels b{{color:#fff;}}
</style>
</head>
<body>
<div class="frame"><i class="tl"></i><i class="tr"></i><i class="bl"></i><i class="br"></i></div>

<div class="juges-wrap">
  <div class="logo-wrap"><img class="logo" src="race4-logo.png" alt="RAC4 · CrossFit Hygie"></div>
  <h1>Planning des juges</h1>
  <p class="intro">Dimanche <b>4 octobre 2026</b>, de <b>7h30 à 13h30</b>, à CrossFit Hygie.<br>
  Trouve ton nom ci-dessous : tu y vois tes heats, ta lane et tes horaires.</p>

  <h2>Ton planning</h2>
  <input class="recherche" id="q" type="search" placeholder="Tape ton nom" autocomplete="off">
  <div class="grille" id="grille">
    {"".join(cartes)}
  </div>

  <h2>Le planning complet</h2>
  <div class="scroll">
    <table>
      <thead><tr><th></th>{th}</tr></thead>
      <tbody>{corps}</tbody>
    </table>
  </div>

  <div class="rappels">
    <b>Sois là 15 minutes avant ton premier heat</b>, briefing juges avant le heat 1.<br>
    Tu juges <b>une lane</b> : tu comptes les répétitions et tu valides les standards de ton athlète.<br>
    Un doute sur un mouvement : tu appelles le <b>head judge</b>, jamais de décision dans le flou.<br>
    Merci d'être là : sans vous, la course n'existe pas 💪
  </div>
</div>

<script>
  var q = document.getElementById('q');
  q.addEventListener('input', function () {{
    var v = q.value.trim().toLowerCase();
    document.querySelectorAll('.juge').forEach(function (c) {{
      c.classList.toggle('cache', v !== '' && c.dataset.nom.indexOf(v) === -1);
    }});
  }});
</script>
</body>
</html>
"""
    OUT.write_text(page, encoding="utf-8")
    print(f"✓ {OUT.name} écrit : {len(juges)} juges, {len(heats)} heats")


if __name__ == "__main__":
    main()
